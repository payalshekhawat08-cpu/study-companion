import hashlib
import json
import logging
import math
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import pymupdf
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from youtube_transcript_api import YouTubeTranscriptApi
from youtube_transcript_api._errors import NoTranscriptFound, TranscriptsDisabled

from config import (
    ENABLE_IMAGE_DESCRIPTIONS,
    MAX_PDF_FIGURES,
    PDF_CHUNK_OVERLAP,
    PDF_CHUNK_SIZE,
    SLIDE_CHUNK_SIZE,
    TOPIC_TAGGING_SNIPPET_SIZE,
    VIDEO_FILE_MAX_SIZE,
    VIDEO_OVERLAP_SECONDS,
    VIDEO_WINDOW_SECONDS,
)
from gemini_client import (
    GeminiError,
    GeminiQuotaError,
    embed_batches,
    generate_json,
    generate_vision_json,
    transcribe_video_file,
)
from prompts import PDF_FIGURE_PROMPT, SLIDE_IMAGE_PROMPT, TOPIC_TAGGING_PROMPT
from store import (
    count_file_chunks,
    existing_chunk_ids,
    get_file_chunks,
    get_image_description,
    get_source,
    get_topic_names,
    get_transcript,
    save_chunks,
    save_image_description,
    save_topic_tags,
    save_transcript,
    record_source_chunk_progress,
    set_source_status,
    set_source_chunk_total,
    set_source_topics,
    update_chunk_tags,
)

BASE_DIR = Path(__file__).resolve().parent
UPLOADS_DIR = BASE_DIR / "uploads"
MAX_FILE_SIZE = 100 * 1024 * 1024
MAX_VIDEO_FILE_SIZE = VIDEO_FILE_MAX_SIZE
logger = logging.getLogger(__name__)


def split_text(
    text: str,
    file_id: str,
    source_file: str,
    source_type: str,
    location: int,
    prefix: str,
    chunk_size: int,
    overlap: int = 0,
):
    clean_text = text.strip()
    chunks = []
    start = 0
    chunk_number = 0
    while start < len(clean_text):
        end = min(start + chunk_size, len(clean_text))
        chunk_text = clean_text[start:end].strip()
        if chunk_text:
            chunks.append(
                {
                    "chunk_id": f"{file_id}:{prefix}:c{chunk_number}",
                    "file_id": file_id,
                    "source_file": source_file,
                    "source_type": source_type,
                    "location": location,
                    "text": chunk_text,
                }
            )
            chunk_number += 1
        if end == len(clean_text):
            break
        start = end - overlap
    return chunks


def read_pdf_chunks(pdf_path: Path, file_id: str, source_file: str):
    chunks = []
    figures = []
    described_count = 0
    with pymupdf.open(pdf_path) as document:
        page_count = document.page_count
        for page_index, page in enumerate(document):
            page_number = page_index + 1
            page_text = page.get_text("text")
            if page_text.strip():
                chunks.extend(
                    split_text(
                        page_text,
                        file_id,
                        source_file,
                        "pdf",
                        page_number,
                        f"p{page_number}",
                        PDF_CHUNK_SIZE,
                        PDF_CHUNK_OVERLAP,
                    )
                )

            page_xrefs = set()
            for image_info in page.get_images(full=True):
                xref, _, width, height = image_info[:4]
                if xref in page_xrefs or width <= 150 or height <= 150:
                    continue
                page_xrefs.add(xref)
                if described_count >= MAX_PDF_FIGURES:
                    break
                image = document.extract_image(xref)
                figure_id = f"{file_id}:p{page_number}:figure{xref}"
                figures.append(
                    (
                        {
                            "chunk_id": figure_id,
                            "file_id": file_id,
                            "source_file": source_file,
                            "source_type": "pdf",
                            "location": page_number,
                        },
                        image["image"],
                        image.get("ext", "png"),
                    )
                )
                described_count += 1
    return page_count, chunks, figures


def read_slide_content(pptx_path: Path, file_id: str, source_file: str):
    presentation = Presentation(str(pptx_path))
    chunks = []
    pictures = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        text_parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text.strip()
                if text:
                    text_parts.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        text_parts.append(row_text)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                picture_id = f"{file_id}:s{slide_number}:image{shape.shape_id}"
                pictures.append(
                    (
                        {
                            "chunk_id": picture_id,
                            "file_id": file_id,
                            "source_file": source_file,
                            "source_type": "slide",
                            "location": slide_number,
                        },
                        shape.image.blob,
                        shape.image.content_type,
                    )
                )

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes and notes not in {"Click to add notes", "Notes"}:
                text_parts.append(notes)

        text = "\n".join(text_parts)
        chunks.extend(
            split_text(
                text,
                file_id,
                source_file,
                "slide",
                slide_number,
                f"s{slide_number}",
                SLIDE_CHUNK_SIZE,
            )
        )
    return len(presentation.slides), chunks, pictures


def _image_mime(extension: str) -> str:
    extension = extension.lower()
    if extension in {"jpg", "jpeg"}:
        return "image/jpeg"
    if extension == "gif":
        return "image/gif"
    if extension == "webp":
        return "image/webp"
    return "image/png"


def _process_pictures(pictures, prompt: str, source_name: str):
    if not ENABLE_IMAGE_DESCRIPTIONS:
        return
    captions = {}
    for picture, image_bytes, mime_type in pictures:
        if existing_chunk_ids([picture["chunk_id"]]):
            continue
        image_hash = hashlib.sha256(image_bytes).hexdigest()
        description = get_image_description(picture["chunk_id"])
        if description is not None:
            captions[image_hash] = description
        elif image_hash not in captions:
            captions[image_hash] = generate_vision_json(
                prompt,
                image=(image_bytes, mime_type),
            )
            save_image_description(picture["chunk_id"], captions[image_hash])
        else:
            save_image_description(picture["chunk_id"], captions[image_hash])
        picture["text"] = picture.get("text_prefix", "") + captions[image_hash]
        _store_chunks([picture], source_name)


def _store_chunks(chunks: list[dict], source_name: str):
    stored_ids = existing_chunk_ids([chunk["chunk_id"] for chunk in chunks])
    missing_chunks = [
        chunk for chunk in chunks if chunk["chunk_id"] not in stored_ids
    ]
    if chunks:
        record_source_chunk_progress(chunks[0]["file_id"])
    missing_texts = [chunk["text"] for chunk in missing_chunks]
    for start, embeddings in embed_batches(
        missing_texts,
        task="document",
        title=source_name,
    ):
        save_chunks([missing_chunks[start]], embeddings)
        record_source_chunk_progress(missing_chunks[start]["file_id"])


def _parse_json(response: str, description: str):
    try:
        parsed = json.loads(response)
    except json.JSONDecodeError as exc:
        raise ValueError(f"Gemini returned invalid JSON for {description}.") from exc
    if not isinstance(parsed, dict):
        raise ValueError(f"Gemini returned invalid JSON for {description}.")
    return parsed


def _tag_source(file_id: str):
    stored_chunks = get_file_chunks(file_id)
    chunks = [
        chunk for chunk in stored_chunks
        if not chunk["metadata"].get("topic")
    ]
    if not chunks:
        set_source_topics(
            file_id,
            sorted(
                {
                    chunk["metadata"]["topic"]
                    for chunk in stored_chunks
                    if chunk["metadata"].get("topic")
                }
            ),
        )
        return
    known_topics = get_topic_names()
    supplied = [
        {
            "number": number,
            "text": " ".join(chunk["text"].split())[:TOPIC_TAGGING_SNIPPET_SIZE],
        }
        for number, chunk in enumerate(chunks, start=1)
    ]
    response = generate_json(
        TOPIC_TAGGING_PROMPT.format(
            existing_topics=json.dumps(known_topics),
            chunks=json.dumps(supplied),
        )
    )
    parsed = _parse_json(response, "topic tags")
    main_topics = parsed.get("topics")
    tagged_chunks = parsed.get("chunks")
    if not isinstance(main_topics, list) or not isinstance(tagged_chunks, list):
        raise ValueError("Gemini returned incomplete topic tags.")
    if any(not isinstance(topic, str) or not topic.strip() for topic in main_topics):
        raise ValueError("Gemini returned an invalid main topic.")

    topic_names = {}
    for topic in main_topics:
        name = topic.strip()
        topic_names.setdefault(name.casefold(), name)
    for existing in known_topics:
        if existing.casefold() in topic_names:
            topic_names[existing.casefold()] = existing
    if not topic_names:
        raise ValueError("Gemini returned no main topics.")

    tags_by_number = {}
    for item in tagged_chunks:
        if not isinstance(item, dict):
            raise ValueError("Gemini returned an invalid chunk tag.")
        number = item.get("number")
        topics = item.get("topics")
        if (
            not isinstance(number, int)
            or isinstance(number, bool)
            or number < 1
            or number > len(chunks)
            or number in tags_by_number
            or not isinstance(topics, list)
            or not 1 <= len(topics) <= 3
            or any(not isinstance(topic, str) for topic in topics)
        ):
            raise ValueError("Gemini returned an invalid chunk tag.")
        normalized_topics = []
        for topic in topics:
            canonical = topic_names.get(topic.strip().casefold())
            if canonical is None:
                raise ValueError("Gemini used a chunk topic missing from main topics.")
            if canonical not in normalized_topics:
                normalized_topics.append(canonical)
        if not normalized_topics:
            raise ValueError("Gemini returned an empty chunk topic list.")
        tags_by_number[number] = normalized_topics

    if set(tags_by_number) != set(range(1, len(chunks) + 1)):
        raise ValueError("Gemini did not tag every chunk.")

    chunk_tags = [
        {
            "chunk_id": chunk["chunk_id"],
            "topic": tags_by_number[number][0],
        }
        for number, chunk in enumerate(chunks, start=1)
    ]
    topic_rows = [
        {"name": topic, "prerequisites": []}
        for topic in topic_names.values()
    ]
    canonical_topics = save_topic_tags(chunk_tags, topic_rows)
    for number, chunk in enumerate(chunks, start=1):
        update_chunk_tags(
            chunk["chunk_id"],
            canonical_topics[chunk["chunk_id"]],
            tags_by_number[number][1:],
        )

    stored_topics = {
        chunk["metadata"]["topic"]
        for chunk in get_file_chunks(file_id)
        if chunk["metadata"].get("topic")
    }
    set_source_topics(
        file_id,
        sorted(stored_topics | set(topic_names.values())),
    )


def _process_transcript(file_id: str, source: dict, transcript: str):
    parsed = _parse_json(transcript, "video transcript")
    segments = parsed.get("segments")
    if not isinstance(segments, list):
        raise ValueError("The video transcript did not contain timestamped segments.")
    clean_segments = []
    for segment in segments:
        if not isinstance(segment, dict):
            continue
        start = segment.get("start")
        text = segment.get("text")
        if (
            isinstance(start, (int, float))
            and math.isfinite(start)
            and start >= 0
            and isinstance(text, str)
            and text.strip()
        ):
            clean_segments.append({"start": float(start), "text": text.strip()})
    if not clean_segments:
        raise ValueError("No transcript captions were found for this video.")
    clean_segments.sort(key=lambda segment: segment["start"])
    last_start = int(clean_segments[-1]["start"])
    stride = VIDEO_WINDOW_SECONDS - VIDEO_OVERLAP_SECONDS
    chunks = []
    for window_start in range(0, last_start + 1, stride):
        window_end = window_start + VIDEO_WINDOW_SECONDS
        window_segments = [
            segment for segment in clean_segments
            if window_start <= segment["start"] < window_end
        ]
        if not window_segments:
            continue
        text = "\n".join(
            f"[{int(segment['start'])}s] {segment['text']}"
            for segment in window_segments
        )
        chunks.append(
            {
                "chunk_id": f"{file_id}:v{window_start}",
                "file_id": file_id,
                "source_file": source["name"],
                "source_type": "video",
                "location": window_start,
                "text": text,
            }
        )
    set_source_chunk_total(file_id, len(chunks))
    _store_chunks(chunks, source["name"])
    _tag_source(file_id)
    return len(chunks)


def _youtube_transcript(video_id: str) -> str:
    try:
        transcript = YouTubeTranscriptApi().fetch(video_id)
    except (NoTranscriptFound, TranscriptsDisabled) as exc:
        raise ValueError("No captions are available for this YouTube video.") from exc
    except Exception as exc:
        raise ValueError(f"The YouTube transcript could not be fetched: {exc}") from exc
    segments = [
        {"start": item.start, "text": item.text}
        for item in transcript
    ]
    return json.dumps({"segments": segments})


def _plain_error(exc: Exception) -> str:
    if isinstance(exc, GeminiQuotaError):
        return "Gemini quota limit reached, wait and try again later."
    if isinstance(exc, GeminiError):
        if exc.error_code == 504:
            return "Gemini took too long to answer. Press retry in a minute."
        stage = exc.stage.lower()
        if "embedding" in stage:
            if exc.error_code == 503:
                return "Gemini is temporarily unavailable while embedding. Use retry."
            return "Internet connection dropped while embedding. Use retry."
        if exc.error_code == 503:
            return "Gemini is temporarily unavailable during processing. Use retry."
        return f"Gemini {stage} failed. Use retry."
    if isinstance(exc, TimeoutError) or "timeout" in type(exc).__name__.lower():
        return "Internet connection timed out during processing. Use retry."
    return "Source processing failed. Use retry."


def process_source(file_id: str):
    source = get_source(file_id)
    if source is None:
        return {"status": "failed", "error": "The uploaded source could not be found."}

    try:
        if source["source_type"] == "video":
            return process_video_source(file_id, source)
        upload_path = UPLOADS_DIR / f"{file_id}{source['storage_ext']}"
        if not upload_path.is_file():
            raise ValueError("The uploaded source file could not be found.")

        if source["source_type"] == "pdf":
            page_count, chunks, figures = read_pdf_chunks(
                upload_path,
                file_id,
                source["name"],
            )
            total_chunks = len(chunks) + (
                len(figures) if ENABLE_IMAGE_DESCRIPTIONS else 0
            )
            set_source_chunk_total(file_id, total_chunks)
            _store_chunks(chunks, source["name"])
            figure_pictures = [
                (
                    figure,
                    image_bytes,
                    _image_mime(extension),
                )
                for figure, image_bytes, extension in figures
            ]
            for figure, image_bytes, mime_type in figure_pictures:
                figure["text_prefix"] = f"[Figure on page {figure['location']}] "
            _process_pictures(figure_pictures, PDF_FIGURE_PROMPT, source["name"])
            _tag_source(file_id)
            set_source_status(file_id, "ready")
            return {
                "status": "ready",
                "pages": page_count,
                "chunks": count_file_chunks(file_id),
            }

        if source["source_type"] == "slide":
            slide_count, chunks, pictures = read_slide_content(
                upload_path,
                file_id,
                source["name"],
            )
            total_chunks = len(chunks) + (
                len(pictures) if ENABLE_IMAGE_DESCRIPTIONS else 0
            )
            set_source_chunk_total(file_id, total_chunks)
            _store_chunks(chunks, source["name"])
            for picture, _, _ in pictures:
                picture["text_prefix"] = f"[Image on slide {picture['location']}] "
            _process_pictures(pictures, SLIDE_IMAGE_PROMPT, source["name"])
            _tag_source(file_id)
            set_source_status(file_id, "ready")
            return {
                "status": "ready",
                "slides": slide_count,
                "chunks": count_file_chunks(file_id),
            }
        raise ValueError("This source type is not supported.")
    except GeminiQuotaError as exc:
        message = _plain_error(exc)
        set_source_status(file_id, "failed", message)
        logger.warning("Ingestion failed for %s: %s", file_id, exc)
        return {"status": "failed", "error": message}
    except Exception as exc:
        message = _plain_error(exc)
        set_source_status(file_id, "failed", message)
        logger.error("Ingestion failed for %s during %s: %s", file_id, getattr(exc, "stage", "processing"), exc)
        return {"status": "failed", "error": message}


def process_video_source(file_id: str, source: dict):
    try:
        transcript = get_transcript(file_id)
        if transcript is None:
            if source["video_url"]:
                parsed_url = urlparse(source["video_url"])
                host = (parsed_url.hostname or "").lower()
                if host == "youtu.be" or host.endswith(".youtu.be"):
                    video_id = parsed_url.path.strip("/").split("/", 1)[0]
                else:
                    video_id = parse_qs(parsed_url.query).get("v", [""])[0]
                    if not video_id:
                        parts = [part for part in parsed_url.path.split("/") if part]
                        video_id = parts[-1] if parts else ""
                transcript = _youtube_transcript(video_id)
            else:
                video_path = UPLOADS_DIR / f"{file_id}{source['storage_ext']}"
                if not video_path.is_file():
                    raise ValueError("The uploaded video file could not be found.")
                transcript = transcribe_video_file(video_path)
            save_transcript(file_id, transcript)

        count = _process_transcript(file_id, source, transcript)
        set_source_status(file_id, "ready")
        return {"status": "ready", "chunks": count_file_chunks(file_id)}
    except GeminiQuotaError as exc:
        message = _plain_error(exc)
        set_source_status(file_id, "failed", message)
        logger.warning("Video ingestion failed for %s: %s", file_id, exc)
        return {"status": "failed", "error": message}
    except Exception as exc:
        message = _plain_error(exc)
        set_source_status(file_id, "failed", message)
        logger.error("Video ingestion failed for %s during %s: %s", file_id, getattr(exc, "stage", "processing"), exc)
        return {"status": "failed", "error": message}
